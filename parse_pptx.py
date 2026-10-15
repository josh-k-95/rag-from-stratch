import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.shapes.picture import Picture


def iter_shapes(shapes):
    """Yield every shape on a slide, descending into grouped shapes."""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes)
        else:
            yield shape


def clean(text):
    # Soft line breaks (Shift+Enter) come back as vertical tabs.
    return text.replace("\v", "\n").strip()


def parse_slide(number, slide, image_dir):
    title_shape = slide.shapes.title
    title = clean(title_shape.text_frame.text) if title_shape is not None else ""
    title_id = title_shape.shape_id if title_shape is not None else None

    # The file stores shapes in creation order, not reading order.
    # Approximate reading order: top-to-bottom, then left-to-right.
    shapes = sorted(iter_shapes(slide.shapes), key=lambda s: (s.top or 0, s.left or 0))

    text_blocks, tables, images = [], [], []
    for shape in shapes:
        if shape.shape_id == title_id:
            continue  # already captured as the title
        if shape.has_table:
            rows = [[clean(cell.text) for cell in row.cells] for row in shape.table.rows]
            tables.append(rows)
        elif isinstance(shape, Picture):
            name = f"slide{number:03d}_{shape.shape_id}.{shape.image.ext}"
            (image_dir / name).write_bytes(shape.image.blob)
            w, h = shape.image.size
            images.append({"file": name, "width_px": w, "height_px": h})
        elif shape.has_text_frame:
            text = clean(shape.text_frame.text)
            if text:
                text_blocks.append(text)

    notes = ""
    if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
        notes = clean(slide.notes_slide.notes_text_frame.text)

    return {
        "slide": number,
        "hidden": slide._element.get("show") == "0",
        "title": title,
        "text_blocks": text_blocks,
        "tables": tables,
        "images": images,
        "notes": notes,
    }


def main():
    if len(sys.argv) != 2:
        sys.exit("usage: python parse_pptx.py data/<deck>.pptx")
    src = Path(sys.argv[1])
    out_dir = Path("parsed")
    image_dir = out_dir / "images"
    image_dir.mkdir(parents=True, exist_ok=True)

    prs = Presentation(src)
    slides = [parse_slide(i, s, image_dir) for i, s in enumerate(prs.slides, start=1)]

    out_file = out_dir / f"{src.stem}_{src.suffix.lstrip(".")}.json"
    out_file.write_text(
        json.dumps({"source_file": src.name, "slides": slides}, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )

    for s in slides:
        hidden = " (hidden)" if s["hidden"] else ""
        empty = not (s["title"] or s["text_blocks"] or s["tables"])
        warn = "   <-- NO EXTRACTABLE TEXT" if empty else ""
        print(
            f"Slide {s['slide']:>3}{hidden} | {s['title'][:40]!r:<42} | "
            f"text:{len(s['text_blocks'])} tables:{len(s['tables'])} "
            f"images:{len(s['images'])} notes:{'y' if s['notes'] else 'n'}{warn}"
        )

    first_table = next((t for s in slides for t in s["tables"]), None)
    if first_table:
        print("\nFirst table found (row by row):")
        for row in first_table:
            print("  ", row)

    n_images = sum(len(s["images"]) for s in slides)
    print(f"\nWrote {out_file} and {n_images} image(s) to {image_dir}/")


if __name__ == "__main__":
    main()